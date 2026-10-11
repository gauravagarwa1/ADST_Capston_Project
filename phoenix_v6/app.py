from flask import Flask, render_template, jsonify, request, redirect, url_for, session, flash, send_from_directory, Response
import sqlite3, hashlib, functools, os, uuid, json, re, mimetypes, csv, io, statistics
from pathlib import Path
from werkzeug.utils import secure_filename

BASE = Path(__file__).resolve().parent
DB = BASE / "data" / "project_phoenix.db"
UPLOADS = BASE / "uploads"
UPLOADS.mkdir(exist_ok=True)
ALLOWED_EXT = {"pdf", "docx", "pptx", "xlsx", "xls", "jpg", "jpeg", "png", "webp", "zip", "csv", "txt"}
MAX_FILE_MB = 50
REPORTS_DIR = BASE / "uploads" / "reports"
REPORTS_DIR.mkdir(parents=True, exist_ok=True)
app = Flask(__name__)
app.secret_key = os.environ.get("PHOENIX_SECRET", "project-phoenix-local-demo-key")


def get_db():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA foreign_keys=ON")
    return c


def rows(items):
    return [dict(x) for x in items]


def hash_password(password):
    return hashlib.sha256(password.encode()).hexdigest()


def login_required(view):
    @functools.wraps(view)
    def wrapper(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapper


def role_required(*roles):
    def deco(view):
        @functools.wraps(view)
        def wrapper(*args, **kwargs):
            if not session.get("user_id"):
                return redirect(url_for("login"))
            if session.get("role") not in roles:
                return redirect(url_for("dashboard"))
            return view(*args, **kwargs)
        return wrapper
    return deco


def _json_safe(value):
    try:
        if value is None: return None
        if hasattr(value, "item"): return value.item()
        if isinstance(value, float) and (value != value or value in (float("inf"), float("-inf"))): return None
        return value
    except Exception:
        return str(value)


def analyze_tabular(path, ext):
    import pandas as pd
    if ext == "csv":
        frame = pd.read_csv(path, nrows=5000)
    else:
        frame = pd.read_excel(path, nrows=5000)
    frame = frame.dropna(axis=1, how="all")
    numeric = frame.select_dtypes(include="number").columns.tolist()
    categorical = [c for c in frame.columns if c not in numeric]
    charts = []
    for c in categorical[:5]:
        counts = frame[c].fillna("(blank)").astype(str).value_counts().head(10)
        charts.append({"type":"bar","title":f"{c} distribution","x":c,"labels":counts.index.tolist(),"values":[_json_safe(x) for x in counts.values.tolist()]})
    for c in numeric[:5]:
        vals = frame[c].dropna()
        if len(vals):
            try:
                counts, edges = __import__('numpy').histogram(vals, bins=min(10, max(4, int(vals.nunique()))))
                labels=[f"{edges[i]:.0f}–{edges[i+1]:.0f}" for i in range(len(counts))]
                values=[int(x) for x in counts.tolist()]
            except Exception:
                labels=[]; values=[]
            charts.append({"type":"hist","title":f"{c} distribution","x":c,"labels":labels,"values":values,"summary":{"min":_json_safe(vals.min()),"max":_json_safe(vals.max()),"average":_json_safe(vals.mean()),"median":_json_safe(vals.median())}})
    if categorical and numeric:
        cat = categorical[0]; num = numeric[0]
        grouped = frame.groupby(cat, dropna=False)[num].mean().sort_values(ascending=False).head(10)
        charts.append({"type":"bar","title":f"Average {num} by {cat}","x":cat,"labels":[str(x) for x in grouped.index],"values":[_json_safe(x) for x in grouped.values.tolist()]})
    missing_by_column = {str(c): int(v) for c,v in frame.isna().sum().items() if int(v)>0}
    duplicate_rows = int(frame.duplicated().sum())
    total_cells = max(1, int(frame.shape[0]*frame.shape[1]))
    missing_cells = int(frame.isna().sum().sum())
    completeness = round(100*(1-missing_cells/total_cells),1)
    unique_by_column = {str(c): int(frame[c].nunique(dropna=True)) for c in frame.columns}
    outliers = {}
    for c in numeric:
        vals=frame[c].dropna()
        if len(vals)>=4:
            q1=float(vals.quantile(.25)); q3=float(vals.quantile(.75)); iqr=q3-q1
            if iqr==0: count=0
            else: count=int(((vals<q1-1.5*iqr)|(vals>q3+1.5*iqr)).sum())
            outliers[str(c)]=count
    numeric_stats=[]
    for c in numeric[:12]:
        vals=frame[c].dropna()
        numeric_stats.append({"column":str(c),"min":_json_safe(vals.min()) if len(vals) else None,"max":_json_safe(vals.max()) if len(vals) else None,"mean":_json_safe(vals.mean()) if len(vals) else None,"median":_json_safe(vals.median()) if len(vals) else None})
    preview = frame.head(30).where(pd.notnull(frame.head(30)), None).to_dict(orient="records")
    summary = {
        "rows": int(len(frame)), "columns": int(len(frame.columns)), "column_names": [str(c) for c in frame.columns],
        "numeric_columns": [str(c) for c in numeric], "categorical_columns": [str(c) for c in categorical],
        "charts": charts, "preview": preview,
        "quality": {"completeness_pct": completeness, "missing_cells": missing_cells, "missing_by_column": missing_by_column,
                    "duplicate_rows": duplicate_rows, "unique_by_column": unique_by_column, "outliers": outliers,
                    "numeric_stats": numeric_stats, "sampled_rows": int(len(frame)), "sample_limit": 5000}
    }
    return summary, frame.to_json(orient="records")

def extract_document(path, ext):
    text = ""
    try:
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:30])
        elif ext == "docx":
            from docx import Document
            doc = Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
        elif ext == "pptx":
            from pptx import Presentation
            prs = Presentation(str(path))
            bits=[]
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        bits.append(shape.text)
            text="\n".join(bits)
        elif ext in {"txt", "csv"}:
            text = path.read_text(errors="ignore")[:200000]
    except Exception as e:
        text = f"Extraction unavailable: {e}"
    return text


def summarize_report(text, ext):
    clean = re.sub(r"\s+", " ", text or "").strip()
    insights=[]
    # Detect common report/chart headings such as the supplied Power BI PDF.
    for line in (text or "").splitlines():
        line=line.strip()
        if not line: continue
        if (" by " in line.lower() or line.lower().startswith(("average", "count", "sum"))) and len(line)<180:
            if line not in insights: insights.append(line)
        if len(insights)>=12: break
    return {"type":ext,"text_preview":clean[:1800],"insights":insights}

def ensure_schema():
    """Small local migration layer so existing V2 databases keep working."""
    conn = get_db()
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS assignments (
        assignment_id INTEGER PRIMARY KEY AUTOINCREMENT,
        title TEXT NOT NULL,
        subject TEXT NOT NULL,
        class_group TEXT NOT NULL,
        description TEXT DEFAULT '',
        skill_level TEXT DEFAULT 'Intermediate',
        due_date TEXT,
        max_points INTEGER DEFAULT 100,
        created_by INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS notifications (
        notification_id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        title TEXT NOT NULL,
        body TEXT NOT NULL,
        is_read INTEGER DEFAULT 0,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(user_id) REFERENCES users(user_id)
    );
    CREATE TABLE IF NOT EXISTS submission_files (
        file_id INTEGER PRIMARY KEY AUTOINCREMENT,
        submission_id INTEGER NOT NULL,
        original_name TEXT NOT NULL,
        stored_name TEXT NOT NULL,
        file_size INTEGER DEFAULT 0,
        uploaded_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(submission_id) REFERENCES submissions(submission_id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS uploaded_reports (
        report_id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        original_name TEXT NOT NULL,
        stored_name TEXT NOT NULL,
        extension TEXT NOT NULL,
        mime_type TEXT DEFAULT '',
        file_size INTEGER DEFAULT 0,
        extracted_text TEXT DEFAULT '',
        analysis_json TEXT DEFAULT '{}',
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(user_id) REFERENCES users(user_id)
    );
    """)
    report_cols = {r[1] for r in conn.execute("PRAGMA table_info(uploaded_reports)").fetchall()}
    for name, definition in {"data_json":"TEXT DEFAULT ''", "row_count":"INTEGER DEFAULT 0", "column_count":"INTEGER DEFAULT 0"}.items():
        if name not in report_cols:
            conn.execute(f"ALTER TABLE uploaded_reports ADD COLUMN {name} {definition}")

    # Add teacher-review fields if the original DB does not have them.
    eval_cols = {r[1] for r in conn.execute("PRAGMA table_info(evaluations)").fetchall()}
    eval_additions = {
        "teacher_score": "REAL",
        "teacher_feedback": "TEXT DEFAULT ''",
        "final_published": "INTEGER DEFAULT 0",
        "reviewed_by": "INTEGER",
        "reviewed_at": "TEXT"
    }
    for name, definition in eval_additions.items():
        if name not in eval_cols:
            conn.execute(f"ALTER TABLE evaluations ADD COLUMN {name} {definition}")

    # V6 migration: these fields belong to the submission itself. The previous
    # version accidentally added them to evaluations, which made project/assignment
    # forms fail because the INSERT targeted columns that did not exist.
    sub_cols = {r[1] for r in conn.execute("PRAGMA table_info(submissions)").fetchall()}
    sub_additions = {
        "assignment_id": "INTEGER",
        "submission_status": "TEXT DEFAULT 'submitted'",
        "reflection": "TEXT DEFAULT ''",
        "project_link": "TEXT DEFAULT ''"
    }
    for name, definition in sub_additions.items():
        if name not in sub_cols:
            conn.execute(f"ALTER TABLE submissions ADD COLUMN {name} {definition}")

    conn.execute("CREATE INDEX IF NOT EXISTS idx_submissions_assignment ON submissions(assignment_id)")
    conn.execute("CREATE INDEX IF NOT EXISTS idx_submissions_student ON submissions(student_id)")
    # Demo assignment records if the table is empty.
    if conn.execute("SELECT COUNT(*) FROM assignments").fetchone()[0] == 0:
        conn.executemany("""INSERT INTO assignments(title,subject,class_group,description,skill_level,due_date,max_points)
            VALUES(?,?,?,?,?,?,?)""", [
            ("Solar-powered smart city", "Physics", "Class 10B", "Design a sustainable neighborhood powered by solar energy.", "Intermediate", "2026-10-30", 100),
            ("Water quality investigation", "Chemistry", "Class 10B", "Investigate local water quality using evidence and clear analysis.", "Intermediate", "2026-10-30", 100),
        ])
    # Bundle the supplied Power BI report as the first dashboard-ready sample.
    sample = REPORTS_DIR / "Project_Phoenix_PowerBI_Report.pdf"
    if sample.exists() and conn.execute("SELECT COUNT(*) FROM uploaded_reports WHERE original_name=?", ("Project_Phoenix_PowerBI_Report.pdf",)).fetchone()[0] == 0:
        try:
            extracted = extract_document(sample, "pdf")
            analysis = summarize_report(extracted, "pdf")
            conn.execute("INSERT INTO uploaded_reports(user_id,original_name,stored_name,extension,mime_type,file_size,extracted_text,analysis_json) VALUES(?,?,?,?,?,?,?,?)",
                (None, sample.name, sample.name, "pdf", "application/pdf", sample.stat().st_size, extracted[:200000], json.dumps(analysis, default=_json_safe)))
        except Exception:
            conn.execute("INSERT INTO uploaded_reports(user_id,original_name,stored_name,extension,mime_type,file_size) VALUES(?,?,?,?,?,?)",
                (None, sample.name, sample.name, "pdf", "application/pdf", sample.stat().st_size))
    conn.commit()
    conn.close()


@app.before_request
def _startup():
    # SQLite is tiny here; this makes a copied project self-healing on first request.
    if DB.exists():
        ensure_schema()


@app.context_processor
def inject_user():
    return {"current_user": session}


@app.route("/")
def home():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))
    return render_template("home.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "")
        role = request.form.get("role", "student")
        conn = get_db()
        u = conn.execute("SELECT * FROM users WHERE username=? AND password_hash=? AND role=?",
                         (username, hash_password(password), role)).fetchone()
        conn.close()
        if u:
            session.clear()
            session.update(user_id=u["user_id"], username=u["username"], role=u["role"],
                           display_name=u["display_name"], student_id=u["student_id"])
            return redirect(url_for("dashboard"))
        flash("Invalid credentials or selected role.", "error")
    return render_template("login.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        role = request.form.get("role", "student")
        name = request.form.get("name", "").strip()
        username = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        if role not in {"student", "teacher", "parent"} or not name or not username or len(password) < 8:
            flash("Please complete the required fields. Password must be 8+ characters.", "error")
            return render_template("register.html")
        conn = get_db()
        try:
            conn.execute("INSERT INTO users(username,password_hash,role,student_id,display_name) VALUES(?,?,?,?,?)",
                         (username, hash_password(password), role, None, name))
            conn.commit()
        except sqlite3.IntegrityError:
            conn.close()
            flash("An account with this email already exists.", "error")
            return render_template("register.html")
        conn.close()
        flash("Account created. You can now sign in.", "success")
        return redirect(url_for("login"))
    return render_template("register.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("home"))


@app.route("/dashboard")
@app.route("/dashboard/")
@login_required
def dashboard():
    return render_template("dashboard.html")


@app.route("/students")
@role_required("teacher")
def students_page():
    return render_template("students.html")


@app.route("/projects")
@login_required
def projects_page():
    return render_template("projects.html")


@app.route("/evaluation")
@role_required("teacher")
def evaluation_page():
    return render_template("evaluation.html")


@app.route("/careers")
@login_required
def careers_page():
    return render_template("careers.html")


@app.route("/leaderboard")
@login_required
def leaderboard_page():
    return render_template("leaderboard.html")


@app.route("/profile")
@login_required
def profile_page():
    return render_template("profile.html")


@app.route("/analytics")
@login_required
def analytics_page():
    return render_template("analytics.html")


@app.route("/settings")
@login_required
def settings_page():
    return render_template("settings.html")


@app.route("/help")
@login_required
def help_page():
    return render_template("help.html")


@app.route("/datasets/<int:report_id>")
@login_required
def dataset_dashboard_page(report_id):
    conn = get_db()
    r = conn.execute("SELECT report_id, original_name, extension, file_size, created_at, analysis_json, stored_name FROM uploaded_reports WHERE report_id=?", (report_id,)).fetchone()
    conn.close()
    if not r:
        return "Dataset not found", 404
    return render_template("dataset_dashboard.html", report=dict(r))


@app.route("/api/settings", methods=["GET", "PUT"])
@login_required
def settings_api():
    conn=get_db()
    if request.method=="PUT":
        d=request.get_json(silent=True) or {}
        name=str(d.get("display_name", "")).strip()
        if name:
            conn.execute("UPDATE users SET display_name=? WHERE user_id=?", (name, session["user_id"]))
            session["display_name"]=name
        conn.commit()
    u=conn.execute("SELECT user_id,username,role,display_name,student_id FROM users WHERE user_id=?", (session["user_id"],)).fetchone()
    conn.close()
    return jsonify(settings=dict(u) if u else {})


@app.route("/reports/<path:filename>")
@login_required
def report_file(filename):
    return send_from_directory(REPORTS_DIR, filename, as_attachment=False)


@app.route("/api/reports", methods=["GET", "POST"])
@login_required
def reports_api():
    conn = get_db()
    if request.method == "POST":
        upload = request.files.get("file")
        if not upload or not upload.filename:
            conn.close(); return jsonify(error="Choose a file first."), 400
        original = secure_filename(upload.filename)
        ext = original.rsplit(".",1)[-1].lower() if "." in original else ""
        if ext not in ALLOWED_EXT:
            conn.close(); return jsonify(error=f"Unsupported format: .{ext}. Supported: PDF, DOCX, PPTX, XLSX, XLS, CSV, TXT and common images/ZIP."), 400
        upload.seek(0, os.SEEK_END); size=upload.tell(); upload.seek(0)
        if size > MAX_FILE_MB*1024*1024:
            conn.close(); return jsonify(error=f"File is larger than {MAX_FILE_MB} MB."), 400
        stored=f"report_{uuid.uuid4().hex}.{ext}"
        target=REPORTS_DIR/stored
        upload.save(target)
        analysis={"type":ext,"text_preview":"","insights":[],"rows":0,"columns":0,"column_names":[],"numeric_columns":[],"categorical_columns":[],"charts":[],"preview":[]}
        extracted=""
        try:
            if ext in {"csv","xlsx","xls"}:
                analysis, _ = analyze_tabular(target, ext)
            elif ext in {"pdf","docx","pptx","txt"}:
                extracted=extract_document(target, ext)
                analysis.update(summarize_report(extracted, ext))
            elif ext in {"jpg","jpeg","png","webp"}:
                analysis.update({"image":True,"text_preview":"Image uploaded successfully. Open the report to view it."})
            else:
                analysis.update({"message":"File stored successfully. Preview is available from the report list."})
        except Exception as e:
            analysis.update({"error":str(e)})
        data_json = ""
        if ext in {"csv","xlsx","xls"}:
            try:
                _, data_json = analyze_tabular(target, ext)
            except Exception:
                data_json = ""
        cur=conn.execute("""INSERT INTO uploaded_reports(user_id,original_name,stored_name,extension,mime_type,file_size,extracted_text,analysis_json,data_json,row_count,column_count) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
            (session["user_id"], original, stored, ext, upload.mimetype or mimetypes.guess_type(original)[0] or "", size, extracted[:200000], json.dumps(analysis, default=_json_safe), data_json, int(analysis.get("rows",0)), int(analysis.get("columns",0))))
        conn.commit(); rid=cur.lastrowid
        conn.close()
        return jsonify(ok=True, report_id=rid, name=original, analysis=analysis, url=url_for("report_file", filename=stored))
    if session.get("role")=="teacher":
        data=conn.execute("SELECT report_id,original_name,extension,file_size,created_at,analysis_json,stored_name FROM uploaded_reports ORDER BY report_id DESC LIMIT 50").fetchall()
    else:
        data=conn.execute("SELECT report_id,original_name,extension,file_size,created_at,analysis_json,stored_name FROM uploaded_reports WHERE user_id=? OR user_id IS NULL ORDER BY report_id DESC LIMIT 50",(session["user_id"],)).fetchall()
    out=[]
    for r in data:
        x=dict(r); x["analysis"]=json.loads(x.pop("analysis_json") or "{}"); x["url"]=url_for("report_file",filename=x.pop("stored_name")); out.append(x)
    conn.close(); return jsonify(reports=out)


@app.route("/api/reports/<int:report_id>/rows")
@login_required
def report_rows(report_id):
    conn=get_db(); r=conn.execute("SELECT * FROM uploaded_reports WHERE report_id=?",(report_id,)).fetchone()
    if not r:
        conn.close(); return jsonify(error="Report not found."),404
    if session.get("role")!="teacher" and r["user_id"] not in (None, session["user_id"]):
        conn.close(); return jsonify(error="Not allowed."),403
    ext=r["extension"]
    if ext not in {"csv","xlsx","xls"}:
        conn.close(); return jsonify(error="This file is not a tabular dataset."),400
    q=request.args.get("q","").strip().lower(); page=max(1,int(request.args.get("page",1))); page_size=min(100,max(5,int(request.args.get("page_size",25))))
    try:
        import pandas as pd
        path=REPORTS_DIR/r["stored_name"]
        frame=pd.read_csv(path,nrows=5000) if ext=="csv" else pd.read_excel(path,nrows=5000)
        frame=frame.dropna(axis=1,how="all")
        if q:
            mask=frame.astype(str).apply(lambda col: col.str.lower().str.contains(q,na=False)).any(axis=1)
            frame=frame[mask]
        total=len(frame); start=(page-1)*page_size; part=frame.iloc[start:start+page_size]
        data=part.where(pd.notnull(part),None).to_dict(orient="records")
        cols=[str(c) for c in frame.columns]
        conn.close(); return jsonify(columns=cols,rows=data,total=total,page=page,page_size=page_size,pages=max(1,(total+page_size-1)//page_size))
    except Exception as e:
        conn.close(); return jsonify(error=str(e)),500

@app.route("/api/reports/<int:report_id>/export")
@login_required
def report_export(report_id):
    conn=get_db(); r=conn.execute("SELECT * FROM uploaded_reports WHERE report_id=?",(report_id,)).fetchone()
    if not r:
        conn.close(); return jsonify(error="Report not found."),404
    if session.get("role")!="teacher" and r["user_id"] not in (None, session["user_id"]):
        conn.close(); return jsonify(error="Not allowed."),403
    if r["extension"] not in {"csv","xlsx","xls"}:
        conn.close(); return jsonify(error="Export is available for CSV/Excel datasets."),400
    try:
        import pandas as pd
        path=REPORTS_DIR/r["stored_name"]
        frame=pd.read_csv(path,nrows=5000) if r["extension"]=="csv" else pd.read_excel(path,nrows=5000)
        out=io.StringIO(); frame.to_csv(out,index=False)
        payload=out.getvalue()
        conn.close()
        return Response(payload,mimetype="text/csv",headers={"Content-Disposition":f'attachment; filename="{Path(r["original_name"]).stem}_export.csv"'})
    except Exception as e:
        conn.close(); return jsonify(error=str(e)),500

@app.route("/api/reports/<int:report_id>", methods=["DELETE"])
@login_required
def delete_report(report_id):
    conn=get_db(); r=conn.execute("SELECT * FROM uploaded_reports WHERE report_id=?",(report_id,)).fetchone()
    if not r:
        conn.close(); return jsonify(error="Report not found."),404
    if session.get("role")!="teacher" and r["user_id"] != session["user_id"]:
        conn.close(); return jsonify(error="Not allowed."),403
    try:
        (REPORTS_DIR/r["stored_name"]).unlink(missing_ok=True)
    except Exception: pass
    conn.execute("DELETE FROM uploaded_reports WHERE report_id=?",(report_id,)); conn.commit(); conn.close()
    return jsonify(ok=True)

@app.route("/api/reports/<int:report_id>")
@login_required
def report_detail(report_id):
    conn=get_db(); r=conn.execute("SELECT * FROM uploaded_reports WHERE report_id=?",(report_id,)).fetchone()
    if not r: conn.close(); return jsonify(error="Report not found."),404
    if session.get("role")!="teacher" and r["user_id"] not in (None, session["user_id"]):
        conn.close(); return jsonify(error="Not allowed."),403
    d=dict(r); d["analysis"]=json.loads(d.pop("analysis_json") or "{}"); d["url"]=url_for("report_file",filename=d.pop("stored_name")); d.pop("extracted_text",None)
    conn.close(); return jsonify(d)

@app.route("/api/dashboard")
@login_required
def dashboard_api():
    conn = get_db()
    role, sid = session.get("role"), session.get("student_id")
    if role in {"student", "parent"} and sid:
        k = conn.execute("""SELECT COUNT(*) total_students, COUNT(p.submission_id) total_submissions,
            ROUND(COALESCE(AVG(p.project_score),0),1) avg_project_score, COALESCE(SUM(p.points_earned),0) total_points,
            ROUND(COALESCE(AVG(p.submitted_on_time),0)*100,1) on_time_pct
            FROM students s LEFT JOIN submissions p ON s.student_id=p.student_id WHERE s.student_id=?""", (sid,)).fetchone()
        subjects = conn.execute("SELECT project_subject label,COUNT(*) value FROM submissions WHERE student_id=? GROUP BY project_subject ORDER BY value DESC", (sid,)).fetchall()
        score = conn.execute("SELECT project_subject label,ROUND(AVG(project_score),1) value FROM submissions WHERE student_id=? GROUP BY project_subject ORDER BY value DESC", (sid,)).fetchall()
    else:
        k = conn.execute("""SELECT COUNT(DISTINCT s.student_id) total_students,
            COUNT(p.submission_id) total_submissions, ROUND(COALESCE(AVG(p.project_score),0),1) avg_project_score,
            COALESCE(SUM(p.points_earned),0) total_points, ROUND(COALESCE(AVG(p.submitted_on_time),0)*100,1) on_time_pct
            FROM students s LEFT JOIN submissions p ON s.student_id=p.student_id""").fetchone()
        subjects = conn.execute("SELECT project_subject label,COUNT(*) value FROM submissions GROUP BY project_subject ORDER BY value DESC").fetchall()
        score = conn.execute("SELECT project_subject label,ROUND(AVG(project_score),1) value FROM submissions GROUP BY project_subject ORDER BY value DESC").fetchall()
    careers = conn.execute("SELECT e.recommended_career label,COUNT(*) value FROM evaluations e WHERE e.recommended_career IS NOT NULL AND e.recommended_career<>'' GROUP BY e.recommended_career ORDER BY value DESC").fetchall()
    education = conn.execute("SELECT education_level label,COUNT(*) value FROM students GROUP BY education_level ORDER BY value DESC").fetchall()
    if role in {"student","parent"} and sid:
        assignments = conn.execute("""SELECT a.*,MAX(p.submission_id) my_submission_id,MAX(p.submission_status) my_submission_status,
            MAX(p.project_score) my_score,MAX(e.teacher_score) my_teacher_score,MAX(e.final_published) my_final_published
            FROM assignments a LEFT JOIN submissions p ON p.assignment_id=a.assignment_id AND p.student_id=?
            LEFT JOIN evaluations e ON e.submission_id=p.submission_id GROUP BY a.assignment_id
            ORDER BY a.due_date ASC,a.assignment_id DESC LIMIT 10""",(sid,)).fetchall()
        recent_feedback = conn.execute("""SELECT p.project_id,p.project_subject,e.teacher_score,e.teacher_feedback,e.reviewed_at
            FROM submissions p JOIN evaluations e ON e.submission_id=p.submission_id
            WHERE p.student_id=? AND e.final_published=1 ORDER BY e.reviewed_at DESC LIMIT 3""",(sid,)).fetchall()
    else:
        assignments = conn.execute("SELECT a.*,COUNT(p.submission_id) submissions FROM assignments a LEFT JOIN submissions p ON p.assignment_id=a.assignment_id GROUP BY a.assignment_id ORDER BY a.due_date ASC,a.assignment_id DESC LIMIT 10").fetchall()
        recent_feedback = conn.execute("""SELECT p.project_id,p.project_subject,s.student_name,e.teacher_score,e.teacher_feedback,e.reviewed_at
            FROM submissions p JOIN students s ON s.student_id=p.student_id JOIN evaluations e ON e.submission_id=p.submission_id
            WHERE e.final_published=1 ORDER BY e.reviewed_at DESC LIMIT 3""").fetchall()
    if sid:
        skill_row = conn.execute("""SELECT ROUND(AVG(programming_score),1) programming,ROUND(AVG(analytical_score),1) analytical,
            ROUND(AVG(creativity_score),1) creativity,ROUND(AVG(communication_score),1) communication,ROUND(AVG(problem_solving_score),1) problem_solving
            FROM submissions WHERE student_id=?""",(sid,)).fetchone()
    else:
        skill_row = conn.execute("""SELECT ROUND(AVG(programming_score),1) programming,ROUND(AVG(analytical_score),1) analytical,
            ROUND(AVG(creativity_score),1) creativity,ROUND(AVG(communication_score),1) communication,ROUND(AVG(problem_solving_score),1) problem_solving
            FROM submissions""").fetchone()
    conn.close()
    return jsonify({"kpis": dict(k), "subjects": rows(subjects), "careers": rows(careers), "education": rows(education), "score": rows(score), "assignments": rows(assignments), "skills": dict(skill_row), "recent_feedback": rows(recent_feedback), "role": role})


@app.route("/api/projects", methods=["GET", "POST"])
@login_required
def projects_api():
    if request.method == "POST":
        d = request.form if request.form else (request.get_json(silent=True) or {})
        role = session.get("role")
        student_id = d.get("student_id") if role == "teacher" else session.get("student_id")
        if not student_id:
            return jsonify(error="A student account is required."), 400
        required = ["project_id", "project_subject", "skill_level", "career_interest"]
        if any(not str(d.get(k, "")).strip() for k in required):
            return jsonify(error="Project title, subject, skill level and career interest are required."), 400
        def num(k, default=0):
            try: return max(0, min(100, float(d.get(k, default))))
            except (TypeError, ValueError): return default
        conn = get_db()
        if not conn.execute("SELECT 1 FROM students WHERE student_id=?", (student_id,)).fetchone():
            conn.close(); return jsonify(error="Student not found."), 404
        assignment_id = d.get("assignment_id") or None
        if assignment_id and not conn.execute("SELECT 1 FROM assignments WHERE assignment_id=?", (assignment_id,)).fetchone():
            conn.close(); return jsonify(error="Selected assignment was not found."), 404
        score = round(num("completeness")*.25 + num("accuracy")*.30 + num("creativity")*.20 + num("problem_solving")*.15 + num("communication")*.10, 1)
        points = int(d.get("points_earned", max(0, round(score))))
        status = "submitted" if str(d.get("submission_status", "submitted")).lower() in {"submitted","submit"} else "draft"
        cur = conn.execute("""INSERT INTO submissions(student_id,project_id,project_subject,skill_level,
            programming_score,analytical_score,creativity_score,communication_score,problem_solving_score,
            completeness_score,accuracy_score,submitted_on_time,project_score,points_earned,career_interest,assignment_id,submission_status,reflection,project_link)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (student_id, str(d["project_id"]).strip(), str(d["project_subject"]).strip(),
            str(d["skill_level"]).strip(), num("programming"), num("analytical"), num("creativity"), num("communication"),
            num("problem_solving"), num("completeness"), num("accuracy"), 1 if str(d.get("submitted_on_time", "1")) in {"1","true","on","yes"} else 0,
            score, points, str(d["career_interest"]).strip(), assignment_id, status, str(d.get("reflection", "")).strip(), str(d.get("project_link", "")).strip()))
        sid_new = cur.lastrowid
        career = recommend_career(num("programming"), num("analytical"), num("creativity"), num("communication"), num("problem_solving"), num("accuracy"))
        feedback = build_feedback(num("completeness"), num("accuracy"), num("creativity"), num("problem_solving"), num("communication"))
        conn.execute("INSERT INTO evaluations(submission_id,ai_feedback,recommended_career) VALUES(?,?,?)", (sid_new, feedback, career))
        for file in request.files.getlist("files"):
            if not file or not file.filename: continue
            ext = file.filename.rsplit(".",1)[-1].lower() if "." in file.filename else ""
            if ext not in ALLOWED_EXT: continue
            file.seek(0, os.SEEK_END); size=file.tell(); file.seek(0)
            if size > MAX_FILE_MB*1024*1024: continue
            stored=f"{sid_new}_{uuid.uuid4().hex}.{ext}"
            file.save(UPLOADS/stored)
            conn.execute("INSERT INTO submission_files(submission_id,original_name,stored_name,file_size) VALUES(?,?,?,?)",(sid_new,file.filename,stored,size))
        # Notify teachers when a student submits a project.
        if status == "submitted":
            teachers = conn.execute("SELECT user_id FROM users WHERE role='teacher'").fetchall()
            for u in teachers:
                conn.execute("INSERT INTO notifications(user_id,title,body) VALUES(?,?,?)", (
                    u[0], "New project submitted", f"{student_id} submitted {d['project_id']} for review."))
        conn.commit(); conn.close()
        return jsonify(ok=True, submission_id=sid_new, project_score=score, recommended_career=career, status=status)

    conn = get_db(); subject=request.args.get("subject","").strip(); skill=request.args.get("skill","").strip()
    where, params=[], []
    if session["role"] in {"student","parent"} and session.get("student_id"):
        where.append("p.student_id=?"); params.append(session["student_id"])
    if subject: where.append("p.project_subject=?"); params.append(subject)
    if skill: where.append("p.skill_level=?"); params.append(skill)
    clause="WHERE "+" AND ".join(where) if where else ""
    data=conn.execute(f"""SELECT p.submission_id,p.project_id,s.student_id,s.student_name,p.project_subject,p.skill_level,
        p.programming_score,p.analytical_score,p.creativity_score,p.communication_score,p.problem_solving_score,
        p.completeness_score,p.accuracy_score,p.points_earned,p.submitted_on_time,p.career_interest,
        p.assignment_id,a.title AS assignment_title,a.due_date,p.submission_status,p.reflection,p.project_link,
        (SELECT COUNT(*) FROM submission_files sf WHERE sf.submission_id=p.submission_id) file_count,
        e.ai_feedback,e.recommended_career,e.teacher_score,e.teacher_feedback,e.final_published
        FROM submissions p JOIN students s ON s.student_id=p.student_id
        LEFT JOIN assignments a ON a.assignment_id=p.assignment_id
        LEFT JOIN evaluations e ON e.submission_id=p.submission_id {clause} ORDER BY p.submission_id DESC LIMIT 500""", params).fetchall()
    subjects=conn.execute("SELECT DISTINCT project_subject FROM submissions ORDER BY project_subject").fetchall()
    skills=conn.execute("SELECT DISTINCT skill_level FROM submissions ORDER BY skill_level").fetchall()
    students=conn.execute("SELECT student_id,student_name FROM students ORDER BY student_name").fetchall() if session["role"]=="teacher" else []
    assignments=conn.execute("SELECT assignment_id,title,subject,class_group,skill_level,due_date,max_points FROM assignments ORDER BY due_date ASC,assignment_id DESC").fetchall()
    conn.close()
    return jsonify({"projects":rows(data),"subjects":[x[0] for x in subjects],"skills":[x[0] for x in skills],"students":rows(students),"assignments":rows(assignments)})


@app.route("/api/assignments", methods=["GET", "POST"])
@login_required
def assignments_api():
    conn = get_db()
    if request.method == "POST":
        if session.get("role") != "teacher":
            conn.close(); return jsonify(error="Only teachers can create assignments."), 403
        d = request.get_json(silent=True) or request.form
        if not all(str(d.get(k, "")).strip() for k in ("title", "subject", "class_group", "due_date")):
            conn.close(); return jsonify(error="Title, subject, class/group and due date are required."), 400
        try:
            max_points = max(1, int(d.get("max_points", 100)))
        except (TypeError, ValueError):
            max_points = 100
        cur = conn.execute("""INSERT INTO assignments(title,subject,class_group,description,skill_level,due_date,max_points,created_by)
            VALUES(?,?,?,?,?,?,?,?)""", (
            str(d["title"]).strip(), str(d["subject"]).strip(), str(d["class_group"]).strip(),
            str(d.get("description", "")).strip(), str(d.get("skill_level", "Intermediate")),
            str(d["due_date"]), max_points, session["user_id"]))
        aid = cur.lastrowid
        # Notify every linked student that a new challenge is available.
        student_users = conn.execute("SELECT user_id FROM users WHERE role='student' AND student_id IS NOT NULL").fetchall()
        for u in student_users:
            conn.execute("INSERT INTO notifications(user_id,title,body) VALUES(?,?,?)", (
                u[0], "New Phoenix assignment", f"{d['title']} · due {d['due_date']}"))
        conn.commit()
        row = conn.execute("SELECT * FROM assignments WHERE assignment_id=?", (aid,)).fetchone()
        conn.close()
        return jsonify(ok=True, assignment_id=aid, assignment=dict(row))

    if session.get("role") in {"student", "parent"} and session.get("student_id"):
        sid = session["student_id"]
        data = conn.execute("""SELECT a.*, COUNT(p.submission_id) AS submissions,
            MAX(p.submission_id) AS my_submission_id,
            MAX(p.submission_status) AS my_submission_status,
            MAX(p.project_score) AS my_score,
            MAX(e.teacher_score) AS my_teacher_score,
            MAX(e.final_published) AS my_final_published
            FROM assignments a
            LEFT JOIN submissions p ON p.assignment_id=a.assignment_id AND p.student_id=?
            LEFT JOIN evaluations e ON e.submission_id=p.submission_id
            GROUP BY a.assignment_id ORDER BY a.due_date ASC, a.assignment_id DESC""", (sid,)).fetchall()
    else:
        data = conn.execute("""SELECT a.*,COUNT(p.submission_id) submissions
            FROM assignments a LEFT JOIN submissions p ON p.assignment_id=a.assignment_id
            GROUP BY a.assignment_id ORDER BY a.due_date ASC,a.assignment_id DESC""").fetchall()
    conn.close()
    return jsonify(assignments=rows(data))


@app.route("/api/students", methods=["GET", "POST"])
@role_required("teacher")
def students_api():
    conn=get_db()
    if request.method=="POST":
        d=request.get_json(silent=True) or request.form
        sid=str(d.get("student_id","")).strip() or generate_student_id(conn)
        name=str(d.get("student_name","")).strip(); city=str(d.get("city","")).strip(); edu=str(d.get("education_level","")).strip()
        if not name or not city or not edu:
            conn.close(); return jsonify(error="Name, city and education level are required."),400
        try:
            conn.execute("INSERT INTO students(student_id,student_name,age,city,education_level) VALUES(?,?,?,?,?)",
                         (sid,name,int(d.get("age",18)),city,edu))
            conn.commit()
        except sqlite3.IntegrityError:
            conn.close(); return jsonify(error="Student ID already exists."),409
        conn.close(); return jsonify(ok=True,student_id=sid)
    q=request.args.get("q","").strip(); city=request.args.get("city","").strip(); education=request.args.get("education","").strip()
    where,params=[],[]
    if q: where.append("(s.student_id LIKE ? OR s.student_name LIKE ?)"); params += [f"%{q}%",f"%{q}%"]
    if city: where.append("s.city=?"); params.append(city)
    if education: where.append("s.education_level=?"); params.append(education)
    clause="WHERE "+" AND ".join(where) if where else ""
    data=conn.execute(f"""SELECT s.student_id,s.student_name,s.age,s.city,s.education_level,
        ROUND(COALESCE(AVG(p.project_score),0),1) avg_score,COALESCE(SUM(p.points_earned),0) points,
        COUNT(p.submission_id) projects,ROUND(COALESCE(AVG(p.submitted_on_time),0)*100,1) on_time_pct
        FROM students s LEFT JOIN submissions p ON s.student_id=p.student_id {clause}
        GROUP BY s.student_id ORDER BY avg_score DESC LIMIT 1000""",params).fetchall()
    cities=conn.execute("SELECT DISTINCT city FROM students ORDER BY city").fetchall(); ed=conn.execute("SELECT DISTINCT education_level FROM students ORDER BY education_level").fetchall()
    conn.close(); return jsonify(students=rows(data),cities=[x[0] for x in cities],educations=[x[0] for x in ed])


def generate_student_id(conn):
    row=conn.execute("SELECT student_id FROM students WHERE student_id LIKE 'STU%' ORDER BY CAST(SUBSTR(student_id,4) AS INTEGER) DESC LIMIT 1").fetchone()
    n=int(row[0][3:])+1 if row else 1
    return f"STU{n:04d}"


def recommend_career(programming, analytical, creativity, communication, problem, accuracy):
    matches=[
        ("Data Analyst", analytical*.55+accuracy*.20+problem*.15+communication*.10),
        ("Software Developer", programming*.55+problem*.20+analytical*.15+communication*.10),
        ("UX / Product Designer", creativity*.55+communication*.20+problem*.15+analytical*.10),
        ("Research Scientist", analytical*.45+accuracy*.25+problem*.20+creativity*.10),
    ]
    return max(matches,key=lambda x:x[1])[0]


def build_feedback(completeness, accuracy, creativity, problem, communication):
    feedback=[]
    if completeness<70: feedback.append("Add missing requirements and supporting evidence.")
    if accuracy<70: feedback.append("Re-check calculations, sources and factual claims.")
    if creativity<70: feedback.append("Explain design choices and try a more original approach.")
    if problem<70: feedback.append("Show your reasoning and how you tested the solution.")
    if communication<70: feedback.append("Improve structure, clarity and visual explanation.")
    return " ".join(feedback) if feedback else "Strong submission profile. Focus on depth, evidence and clear reflection."


@app.route("/api/evaluate", methods=["POST"])
@role_required("teacher")
def evaluate_api():
    d=request.get_json(force=True)
    try: submission_id=int(d.get("submission_id"))
    except (TypeError, ValueError): return jsonify(error="Submission ID is required."),400
    def n(k):
        try:return max(0,min(100,float(d.get(k,0))))
        except:return 0
    completeness,accuracy,creativity,problem,communication=n("completeness"),n("accuracy"),n("creativity"),n("problem_solving"),n("communication")
    score=round(completeness*.25+accuracy*.30+creativity*.20+problem*.15+communication*.10,1)
    career=recommend_career(n("programming"),n("analytical"),creativity,communication,problem,accuracy)
    feedback=build_feedback(completeness,accuracy,creativity,problem,communication)
    matches=[
        ("Data Analyst",n("analytical")*.55+accuracy*.20+problem*.15+communication*.10),
        ("Software Developer",n("programming")*.55+problem*.20+n("analytical")*.15+communication*.10),
        ("UX / Product Designer",creativity*.55+communication*.20+problem*.15+n("analytical")*.10),
        ("Research Scientist",n("analytical")*.45+accuracy*.25+problem*.20+creativity*.10),
    ]
    conn=get_db()
    if not conn.execute("SELECT 1 FROM submissions WHERE submission_id=?",(submission_id,)).fetchone():
        conn.close(); return jsonify(error="Submission not found."),404
    conn.execute("""UPDATE submissions SET programming_score=?,analytical_score=?,creativity_score=?,communication_score=?,
        problem_solving_score=?,completeness_score=?,accuracy_score=?,project_score=? WHERE submission_id=?""",
        (n("programming"),n("analytical"),creativity,communication,problem,completeness,accuracy,score,submission_id))
    conn.execute("UPDATE evaluations SET ai_feedback=?,recommended_career=? WHERE submission_id=?",(feedback,career,submission_id))
    conn.commit(); conn.close()
    return jsonify(project_score=score,recommended_career=career,feedback=feedback,career_matches=[{"career":a,"score":round(b,1)} for a,b in sorted(matches,key=lambda x:x[1],reverse=True)])


@app.route("/api/evaluation/final", methods=["POST"])
@role_required("teacher")
def final_evaluation_api():
    d=request.get_json(force=True); submission_id=d.get("submission_id")
    try: score=max(0,min(100,float(d.get("teacher_score"))))
    except: return jsonify(error="Teacher score must be a number from 0 to 100."),400
    feedback=str(d.get("teacher_feedback","")).strip()
    if not submission_id:return jsonify(error="Submission ID is required."),400
    conn=get_db(); exists=conn.execute("SELECT submission_id FROM submissions WHERE submission_id=?",(submission_id,)).fetchone()
    if not exists:conn.close();return jsonify(error="Submission not found."),404
    points=round(score)
    conn.execute("""UPDATE evaluations SET teacher_score=?,teacher_feedback=?,final_published=?,reviewed_by=?,reviewed_at=CURRENT_TIMESTAMP WHERE submission_id=?""",
                 (score,feedback,1 if d.get("publish") else 0,session["user_id"],submission_id))
    if d.get("publish"):
        student = conn.execute("SELECT student_id,project_id FROM submissions WHERE submission_id=?",(submission_id,)).fetchone()
        conn.execute("UPDATE submissions SET project_score=?,points_earned=?,submission_status='reviewed' WHERE submission_id=?",(score,points,submission_id))
        if student:
            users = conn.execute("SELECT user_id FROM users WHERE student_id=? AND role IN ('student','parent')",(student["student_id"],)).fetchall()
            for u in users:
                conn.execute("INSERT INTO notifications(user_id,title,body) VALUES(?,?,?)",(u[0],"Final grade published",f"{student['project_id']} received {score:.0f}/100."))
    conn.commit();conn.close();return jsonify(ok=True,final_published=bool(d.get("publish")),score=score)


@app.route("/api/careers")
@login_required
def careers_api():
    conn=get_db(); sid=session.get("student_id") if session["role"] in {"student","parent"} else None
    where="WHERE p.student_id=?" if sid else ""; params=[sid] if sid else []
    dist=conn.execute(f"SELECT e.recommended_career label,COUNT(*) value FROM submissions p JOIN evaluations e ON e.submission_id=p.submission_id {where} GROUP BY e.recommended_career ORDER BY value DESC",params).fetchall()
    matrix=conn.execute(f"SELECT p.career_interest interest,e.recommended_career career,COUNT(*) value FROM submissions p JOIN evaluations e ON e.submission_id=p.submission_id {where} GROUP BY p.career_interest,e.recommended_career ORDER BY value DESC",params).fetchall()
    conn.close();return jsonify(distribution=rows(dist),interest_matrix=rows(matrix))


@app.route("/api/leaderboard")
@login_required
def leaderboard_api():
    conn=get_db();data=conn.execute("""SELECT s.student_id,s.student_name,s.city,s.education_level,
        COALESCE(SUM(p.points_earned),0) points,ROUND(COALESCE(AVG(p.project_score),0),1) avg_score,COUNT(p.submission_id) projects
        FROM students s JOIN submissions p ON s.student_id=p.student_id GROUP BY s.student_id ORDER BY points DESC,avg_score DESC LIMIT 50""").fetchall();conn.close();return jsonify(leaderboard=rows(data))


@app.route("/api/profile", methods=["GET","PUT"])
@login_required
def profile_api():
    conn=get_db();sid=session.get("student_id")
    if request.method=="PUT":
        if not sid:conn.close();return jsonify(error="This account is not linked to a student."),400
        d=request.get_json(silent=True) or {}
        name=str(d.get("student_name","")).strip(); city=str(d.get("city","")).strip(); edu=str(d.get("education_level","")).strip()
        conn.execute("UPDATE students SET student_name=COALESCE(NULLIF(?,''),student_name),city=COALESCE(NULLIF(?,''),city),education_level=COALESCE(NULLIF(?,''),education_level) WHERE student_id=?",(name,city,edu,sid))
        if name: conn.execute("UPDATE users SET display_name=? WHERE user_id=?",(name,session["user_id"])); session["display_name"]=name
        conn.commit();conn.close();return jsonify(ok=True)
    if sid:st=conn.execute("SELECT * FROM students WHERE student_id=?",(sid,)).fetchone();projects=conn.execute("""SELECT p.submission_id,p.project_id,p.project_subject,p.project_score,p.points_earned,e.recommended_career,e.ai_feedback,e.final_published FROM submissions p LEFT JOIN evaluations e ON e.submission_id=p.submission_id WHERE p.student_id=? ORDER BY p.project_score DESC""",(sid,)).fetchall()
    else:st=conn.execute("SELECT * FROM students ORDER BY student_id LIMIT 1").fetchone();projects=[]
    conn.close();return jsonify(student=dict(st) if st else {},projects=rows(projects))


@app.route("/api/submissions/<int:submission_id>")
@login_required
def submission_detail_api(submission_id):
    conn=get_db()
    row=conn.execute("""SELECT p.*,s.student_name,a.title assignment_title,a.description assignment_description,a.due_date,a.max_points,
        e.ai_feedback,e.recommended_career,e.teacher_score,e.teacher_feedback,e.final_published
        FROM submissions p JOIN students s ON s.student_id=p.student_id
        LEFT JOIN assignments a ON a.assignment_id=p.assignment_id
        LEFT JOIN evaluations e ON e.submission_id=p.submission_id
        WHERE p.submission_id=?""",(submission_id,)).fetchone()
    if not row:
        conn.close(); return jsonify(error="Submission not found."),404
    if session.get("role") in {"student","parent"} and row["student_id"] != session.get("student_id"):
        conn.close(); return jsonify(error="Not allowed."),403
    conn.close(); return jsonify(submission=dict(row))

@app.route("/uploads/<path:name>")
@login_required
def uploaded_file(name):
    # Only authenticated users can open uploaded project files.
    return send_from_directory(UPLOADS, name, as_attachment=False)

@app.route("/api/submission/<int:submission_id>/files")
@login_required
def submission_files_api(submission_id):
    conn=get_db()
    row=conn.execute("SELECT student_id FROM submissions WHERE submission_id=?",(submission_id,)).fetchone()
    if not row:
        conn.close(); return jsonify(error="Submission not found."),404
    if session.get("role") in {"student","parent"} and row["student_id"] != session.get("student_id"):
        conn.close(); return jsonify(error="Not allowed."),403
    files=conn.execute("SELECT file_id,original_name,file_size,uploaded_at,stored_name FROM submission_files WHERE submission_id=? ORDER BY file_id",(submission_id,)).fetchall()
    conn.close(); return jsonify(files=[dict(x) | {"url":url_for("uploaded_file",name=x["stored_name"])} for x in files])

@app.route("/notifications")
@login_required
def notifications_page():
    return render_template("notifications.html")

@app.route("/api/notifications", methods=["GET","POST"])
@login_required
def notifications_api():
    conn=get_db()
    if request.method=="POST":
        d=request.get_json(silent=True) or {}
        if d.get("all"):
            conn.execute("UPDATE notifications SET is_read=1 WHERE user_id=?",(session["user_id"],))
        elif d.get("notification_id"):
            conn.execute("UPDATE notifications SET is_read=1 WHERE notification_id=? AND user_id=?",(d["notification_id"],session["user_id"]))
        conn.commit()
    data=conn.execute("SELECT notification_id,title,body,is_read,created_at FROM notifications WHERE user_id=? ORDER BY notification_id DESC LIMIT 100",(session["user_id"],)).fetchall()
    unread=conn.execute("SELECT COUNT(*) FROM notifications WHERE user_id=? AND is_read=0",(session["user_id"],)).fetchone()[0]
    conn.close(); return jsonify(notifications=rows(data),unread=unread)

@app.route("/api/analytics")
@login_required
def analytics_api():
    conn=get_db(); role=session.get("role"); sid=session.get("student_id")
    where="WHERE p.student_id=?" if role in {"student","parent"} and sid else ""; params=[sid] if where else []
    # Dataset-backed metrics. Every graph is computed from the current SQLite state, so newly added data appears automatically.
    k=conn.execute(f"""SELECT COUNT(DISTINCT p.student_id) students,COUNT(p.submission_id) projects,
        ROUND(COALESCE(AVG(p.project_score),0),1) avg_score,COALESCE(SUM(p.points_earned),0) points,
        ROUND(COALESCE(AVG(p.accuracy_score),0),1) avg_accuracy,ROUND(COALESCE(AVG(p.completeness_score),0),1) avg_completeness
        FROM submissions p {where}""",params).fetchone()
    subject=conn.execute(f"SELECT p.project_subject label,COUNT(*) value FROM submissions p {where} GROUP BY p.project_subject ORDER BY value DESC",params).fetchall()
    subject_score=conn.execute(f"SELECT p.project_subject label,ROUND(AVG(p.project_score),1) value FROM submissions p {where} GROUP BY p.project_subject ORDER BY value DESC",params).fetchall()
    skill=conn.execute(f"SELECT p.skill_level label,COUNT(*) value FROM submissions p {where} GROUP BY p.skill_level ORDER BY value DESC",params).fetchall()
    careers=conn.execute(f"SELECT e.recommended_career label,COUNT(*) value FROM submissions p LEFT JOIN evaluations e ON e.submission_id=p.submission_id {where} GROUP BY e.recommended_career ORDER BY value DESC",params).fetchall()
    education=conn.execute("SELECT education_level label,COUNT(*) value FROM students GROUP BY education_level ORDER BY value DESC").fetchall() if not where else conn.execute("SELECT s.education_level label,COUNT(*) value FROM students s WHERE s.student_id=? GROUP BY s.education_level",params).fetchall()
    scores=conn.execute(f"SELECT p.completeness_score x,p.accuracy_score y,p.project_score size,p.project_subject label FROM submissions p {where} ORDER BY p.submission_id",params).fetchall()
    skills=conn.execute(f"""SELECT ROUND(AVG(programming_score),1) programming,ROUND(AVG(analytical_score),1) analytical,
        ROUND(AVG(creativity_score),1) creativity,ROUND(AVG(communication_score),1) communication,
        ROUND(AVG(problem_solving_score),1) problem_solving FROM submissions p {where}""",params).fetchone()
    ontime=conn.execute(f"SELECT CASE WHEN p.submitted_on_time=1 THEN 'On time' ELSE 'Late' END label,COUNT(*) value FROM submissions p {where} GROUP BY p.submitted_on_time",params).fetchall()
    points_subject=conn.execute(f"SELECT p.project_subject label,COALESCE(SUM(p.points_earned),0) value FROM submissions p {where} GROUP BY p.project_subject ORDER BY value DESC",params).fetchall()
    assignment_subject=conn.execute("SELECT subject label,COUNT(*) value FROM assignments GROUP BY subject ORDER BY value DESC").fetchall()
    if where:
        top=conn.execute("SELECT s.student_name label,COALESCE(SUM(p.points_earned),0) value FROM students s JOIN submissions p ON p.student_id=s.student_id WHERE s.student_id=? GROUP BY s.student_id",params).fetchall()
    else:
        top=conn.execute("SELECT s.student_name label,COALESCE(SUM(p.points_earned),0) value FROM students s JOIN submissions p ON p.student_id=s.student_id GROUP BY s.student_id ORDER BY value DESC LIMIT 10").fetchall()
    status=conn.execute(f"SELECT COALESCE(p.submission_status,'submitted') label,COUNT(*) value FROM submissions p {where} GROUP BY COALESCE(p.submission_status,'submitted')",params).fetchall()
    career_matrix=conn.execute(f"SELECT p.career_interest interest,e.recommended_career career,COUNT(*) value FROM submissions p LEFT JOIN evaluations e ON e.submission_id=p.submission_id {where} GROUP BY p.career_interest,e.recommended_career ORDER BY value DESC",params).fetchall()
    conn.close()
    return jsonify(kpis=dict(k),subject=rows(subject),subject_score=rows(subject_score),skill=rows(skill),careers=rows(careers),education=rows(education),accuracy_completeness=rows(scores),skills=dict(skills),ontime=rows(ontime),points_subject=rows(points_subject),assignment_subject=rows(assignment_subject),top_students=rows(top),status=rows(status),career_matrix=rows(career_matrix))


if __name__ == "__main__":
    ensure_schema()
    app.run(debug=True, host="127.0.0.1", port=5000)
